from pathlib import Path
from typing  import Iterable
from langchain_core .documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader,TextLoader
from docx import Document as DocxDocument


SUPPORTED={
    ".pdf",".txt",".docx",".md"
}

# loading our input data from internal or external sources
def load_file(path:Path)->list[Document]:
    suffix=path.suffix.lower()
    if suffix==".pdf":
        return PyPDFLoader(str(path)).load()
    if suffix in {".txt",".md"}:
        return TextLoader(str(path),encoding="utf-8").load()
    if suffix==".docx":
        doc=DocxDocument(str(path))
        text="\n".join([para.text for para in doc.paragraphs])
        return [Document(page_content=text,metadata={"source":str(path)})]
    raise ValueError(f"Unsupported file type: {suffix}")


# chinking the documents into smaller pieces


def chunk_documents(docs:Iterable[Document],chunk_size:int=1000,chunk_overlap:int=200)->list[Document]:
    text_splitter=RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        length_function=len
    )
    return text_splitter.split_documents(docs)

